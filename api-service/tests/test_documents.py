import io
import zipfile

import pytest
from docx import Document
from pptx import Presentation

from classfox.documents import DocumentError, parse_document


def test_extracts_docx_paragraphs_and_tables() -> None:
    document = Document()
    document.add_paragraph("数据库系统课程")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "事务"
    table.cell(0, 1).text = "原子性"
    buffer = io.BytesIO()
    document.save(buffer)
    text = parse_document(buffer.getvalue(), "课程.docx")
    assert "数据库系统课程" in text
    assert "事务 | 原子性" in text


def test_extracts_pptx_and_rejects_empty_presentations() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "内存管理"
    buffer = io.BytesIO()
    presentation.save(buffer)
    assert "内存管理" in parse_document(buffer.getvalue(), "lesson.pptx")
    empty = io.BytesIO()
    Presentation().save(empty)
    with pytest.raises(DocumentError, match="文字"):
        parse_document(empty.getvalue(), "empty.pptx")


@pytest.mark.parametrize("filename", ["legacy.doc", "legacy.ppt", "script.html", "archive.zip"])
def test_unsupported_formats_have_an_actionable_error(filename: str) -> None:
    with pytest.raises(DocumentError, match="格式"):
        parse_document(b"synthetic content", filename)


def test_corrupt_file_does_not_leak_parser_internals() -> None:
    with pytest.raises(DocumentError, match="损坏"):
        parse_document(b"not a real office file", "document.docx")


def test_compressed_archive_resource_limit_is_enforced_before_parsing() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "x" * 2_000_000)
    with pytest.raises(DocumentError, match="解压"):
        parse_document(buffer.getvalue(), "compressed.docx")
