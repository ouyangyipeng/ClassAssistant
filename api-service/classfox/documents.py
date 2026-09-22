import io
import zipfile
from pathlib import PurePath

from docx import Document
from pptx import Presentation
from pypdf import PdfReader

MAX_TEXT_LENGTH = 2_000_000
MAX_FILE_SIZE = 25 * 1024 * 1024
SUPPORTED_FORMATS = {".pdf", ".pptx", ".docx", ".txt", ".md"}


class DocumentError(ValueError):
    """An actionable parsing failure safe to show to users."""


def check_archive(data: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        unpacked = sum(entry.file_size for entry in entries)
        if len(entries) > 10000 or unpacked > 128 * 1024 * 1024:
            raise DocumentError("资料解压后体积过大，请拆分后上传")
        if any(entry.file_size > 1_000_000 and entry.file_size > max(entry.compress_size, 1) * 200 for entry in entries):
            raise DocumentError("资料解压比例异常，请重新导出文件")


def extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise DocumentError("请先移除 PDF 密码保护后上传")
    if len(reader.pages) > 1000:
        raise DocumentError("PDF 超过 1000 页，请拆分后上传")
    sections: list[str] = []
    total = 0
    for index, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        total += len(text)
        if total > MAX_TEXT_LENGTH:
            raise DocumentError("资料文字过多，请拆分后上传")
        if text:
            sections.append(f"第 {index} 页\n{text}")
    return "\n\n".join(sections)


def extract_docx(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    lines = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        lines.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
    return "\n".join(lines)


def extract_pptx(data: bytes) -> str:
    presentation = Presentation(io.BytesIO(data))
    if len(presentation.slides) > 1000:
        raise DocumentError("课件超过 1000 页，请拆分后上传")
    sections: list[str] = []
    for index, slide in enumerate(presentation.slides, 1):
        lines: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines.append(shape.text_frame.text.strip())
            if shape.has_table:
                lines.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in shape.table.rows)
        if any(lines):
            sections.append(f"第 {index} 页\n" + "\n".join(lines))
    return "\n\n".join(sections)


def parse_document(data: bytes, filename: str) -> str:
    extension = PurePath(filename).suffix.lower()
    if extension not in SUPPORTED_FORMATS:
        raise DocumentError("不支持的资料格式，请将旧 .doc/.ppt 转为 .docx/.pptx，或上传 PDF/TXT/Markdown")
    if len(data) > MAX_FILE_SIZE:
        raise DocumentError("资料超过 25 MiB，请拆分后上传")
    try:
        if extension in {".docx", ".pptx"}:
            check_archive(data)
        if extension == ".pdf":
            text = extract_pdf(data)
        elif extension == ".docx":
            text = extract_docx(data)
        elif extension == ".pptx":
            text = extract_pptx(data)
        else:
            text = data.decode("utf-8-sig")
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError("资料损坏或格式不正确，请重新导出后上传") from exc
    if not text.strip():
        raise DocumentError("资料没有可提取文字；扫描 PDF 请先完成 OCR")
    if len(text) > MAX_TEXT_LENGTH:
        raise DocumentError("资料文字过多，请拆分后上传")
    return text.strip()
